import csv
import hashlib
import logging
import pickle
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from app.embeddings import (
    cosine_similarity,
    embed_text,
)


# =========================================================
# CONFIGURATION
# =========================================================

logger = logging.getLogger(__name__)

KNOWLEDGE_DIR = Path("knowledge")
INDEX_PATH = Path("data") / "vector_index.pkl"

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".xlsx",
    ".xlsm",
    ".pptx",
    ".pdf",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
}


# =========================================================
# TXT LOADER
# =========================================================

def load_txt_document(
    file_path: Path,
) -> str:
    """
    Read a plain-text knowledge document.
    """

    return file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


# =========================================================
# EXCEL LOADER
# =========================================================

def load_excel_document(
    file_path: Path,
) -> str:
    """
    Convert an Excel workbook into searchable text.

    Each sheet and row is preserved in readable form so that
    cutter benchmarking data can be retrieved by the RAG layer.
    """

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    sections: list[str] = []

    try:

        for worksheet in workbook.worksheets:

            sections.append(
                f"WORKSHEET: {worksheet.title}"
            )

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):

                values = []

                for value in row:

                    if value is None:
                        continue

                    cleaned_value = str(value).strip()

                    if cleaned_value:
                        values.append(cleaned_value)

                if not values:
                    continue

                row_text = " | ".join(values)

                sections.append(
                    f"Row {row_number}: {row_text}"
                )

            sections.append("")

    finally:

        workbook.close()

    return "\n".join(sections).strip()


# =========================================================
# POWERPOINT HELPERS
# =========================================================

def extract_powerpoint_shape_text(
    shape,
) -> list[str]:
    """
    Extract text from normal PowerPoint shapes,
    tables, and grouped shapes.
    """

    extracted_text: list[str] = []

    # -----------------------------------------------------
    # Normal text boxes / titles / placeholders
    # -----------------------------------------------------

    if getattr(
        shape,
        "has_text_frame",
        False,
    ):

        text = shape.text.strip()

        if text:
            extracted_text.append(text)

    # -----------------------------------------------------
    # Tables
    # -----------------------------------------------------

    if getattr(
        shape,
        "has_table",
        False,
    ):

        table = shape.table

        for row_number, row in enumerate(
            table.rows,
            start=1,
        ):

            values = []

            for cell in row.cells:

                value = cell.text.strip()

                if value:
                    values.append(value)

            if values:

                extracted_text.append(
                    f"Table Row {row_number}: "
                    + " | ".join(values)
                )

    # -----------------------------------------------------
    # Grouped shapes
    # -----------------------------------------------------

    if hasattr(
        shape,
        "shapes",
    ):

        for child_shape in shape.shapes:

            extracted_text.extend(
                extract_powerpoint_shape_text(
                    child_shape
                )
            )

    return extracted_text


# =========================================================
# POWERPOINT LOADER
# =========================================================

def load_powerpoint_document(
    file_path: Path,
) -> str:
    """
    Convert PowerPoint slides into searchable text.
    """

    presentation = Presentation(
        file_path
    )

    sections: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        sections.append(
            f"SLIDE: {slide_number}"
        )

        slide_text: list[str] = []

        for shape in slide.shapes:

            slide_text.extend(
                extract_powerpoint_shape_text(
                    shape
                )
            )

        if slide_text:

            sections.extend(
                slide_text
            )

        sections.append("")

    return "\n".join(sections).strip()


# =========================================================
# CSV LOADER
# =========================================================

def load_csv_document(
    file_path: Path,
) -> str:
    """
    Convert a CSV file into searchable text.
    """

    sections: list[str] = []

    with file_path.open(
        "r",
        encoding="utf-8",
        errors="ignore",
        newline="",
    ) as file:

        reader = csv.reader(file)

        for row_number, row in enumerate(
            reader,
            start=1,
        ):

            values = [
                cell.strip()
                for cell in row
                if cell.strip()
            ]

            if not values:
                continue

            sections.append(
                f"Row {row_number}: "
                + " | ".join(values)
            )

    return "\n".join(sections).strip()


# =========================================================
# PDF LOADER
# =========================================================

def load_pdf_document(
    file_path: Path,
) -> str:
    """
    Extract text from a PDF document, page by page.
    """

    from pypdf import PdfReader

    reader = PdfReader(
        str(file_path)
    )

    sections: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        text = (
            page.extract_text()
            or ""
        ).strip()

        if not text:
            continue

        sections.append(
            f"PAGE: {page_number}"
        )
        sections.append(text)
        sections.append("")

    return "\n".join(sections).strip()


# =========================================================
# IMAGE LOADER (OCR)
# =========================================================

def load_image_document(
    file_path: Path,
) -> str:
    """
    Extract visible text from an image using OCR.

    Requires the Tesseract OCR engine to be installed
    and available on the system PATH. If it is not
    available, the image is skipped rather than
    failing the whole knowledge load.
    """

    import pytesseract
    from PIL import Image

    with Image.open(file_path) as image:

        text = pytesseract.image_to_string(
            image
        )

    return text.strip()


# =========================================================
# GENERIC DOCUMENT LOADER
# =========================================================

def load_document(
    file_path: Path,
) -> str:
    """
    Load a supported knowledge document and normalize
    its contents into plain searchable text.
    """

    suffix = file_path.suffix.lower()

    if suffix == ".txt":

        return load_txt_document(
            file_path
        )

    if suffix in {
        ".xlsx",
        ".xlsm",
    }:

        return load_excel_document(
            file_path
        )

    if suffix == ".pptx":

        return load_powerpoint_document(
            file_path
        )

    if suffix == ".csv":

        return load_csv_document(
            file_path
        )

    if suffix == ".pdf":

        return load_pdf_document(
            file_path
        )

    if suffix in {
        ".png",
        ".jpg",
        ".jpeg",
    }:

        return load_image_document(
            file_path
        )

    return ""


# =========================================================
# LOAD ALL KNOWLEDGE DOCUMENTS
# =========================================================

def load_knowledge_documents() -> list[dict]:
    """
    Load every supported document from the knowledge folder.

    Supported:
        .txt
        .xlsx / .xlsm
        .pptx
        .pdf
        .csv
        .png / .jpg / .jpeg (OCR)
    """

    documents = []

    if not KNOWLEDGE_DIR.exists():

        logger.warning(
            "Knowledge directory does not exist: %s",
            KNOWLEDGE_DIR,
        )

        return documents

    for file_path in sorted(
        KNOWLEDGE_DIR.iterdir(),
        key=lambda path: path.name.lower(),
    ):

        # Skip directories
        if not file_path.is_file():
            continue

        # Skip temporary Microsoft Office files
        if file_path.name.startswith("~$"):
            continue

        # Skip unsupported formats
        if (
            file_path.suffix.lower()
            not in SUPPORTED_EXTENSIONS
        ):
            continue

        try:

            text = load_document(
                file_path
            )

        except Exception as exc:

            logger.exception(
                "Failed to load knowledge document %s: %s",
                file_path.name,
                exc,
            )

            continue

        text = text.strip()

        if not text:

            logger.warning(
                "Knowledge document contains no readable text: %s",
                file_path.name,
            )

            continue

        documents.append(
            {
                "source": file_path.name,
                "content": text,
            }
        )

        logger.info(
            "Knowledge document loaded: %s",
            file_path.name,
        )

    logger.info(
        "Knowledge documents loaded count=%s",
        len(documents),
    )

    return documents


# =========================================================
# CHUNK TEXT
# =========================================================

def chunk_text(
    text: str,
    chunk_size: int = 500,
    overlap: int = 100,
) -> list[str]:

    if chunk_size <= 0:

        raise ValueError(
            "chunk_size must be greater than 0."
        )

    if overlap < 0:

        raise ValueError(
            "overlap cannot be negative."
        )

    if overlap >= chunk_size:

        raise ValueError(
            "overlap must be smaller than chunk_size."
        )

    cleaned_text = text.strip()

    if not cleaned_text:
        return []

    chunks = []

    start = 0

    while start < len(
        cleaned_text
    ):

        end = (
            start
            + chunk_size
        )

        chunk = cleaned_text[
            start:end
        ].strip()

        if chunk:

            chunks.append(
                chunk
            )

        if end >= len(
            cleaned_text
        ):

            break

        start = (
            end
            - overlap
        )

    return chunks


# =========================================================
# CREATE KNOWLEDGE CHUNKS
# =========================================================

def create_knowledge_chunks(
    chunk_size: int = 500,
    overlap: int = 100,
) -> list[dict]:

    documents = (
        load_knowledge_documents()
    )

    chunks = []

    for document in documents:

        document_chunks = (
            chunk_text(
                text=document[
                    "content"
                ],
                chunk_size=chunk_size,
                overlap=overlap,
            )
        )

        for (
            chunk_index,
            chunk,
        ) in enumerate(
            document_chunks
        ):

            chunks.append(
                {
                    "source": (
                        document[
                            "source"
                        ]
                    ),
                    "chunk_id": (
                        f"{document['source']}"
                        f"::chunk-{chunk_index}"
                    ),
                    "chunk_index": (
                        chunk_index
                    ),
                    "content": (
                        chunk
                    ),
                }
            )

    logger.info(
        "Knowledge chunks created count=%s",
        len(chunks),
    )

    return chunks


# =========================================================
# KNOWLEDGE FINGERPRINT
# =========================================================

def calculate_knowledge_fingerprint() -> str:
    """
    Generate a fingerprint representing all currently
    loaded knowledge documents.

    If any document changes or a new document is added,
    the fingerprint changes and the vector index rebuilds.
    """

    hasher = hashlib.sha256()

    documents = (
        load_knowledge_documents()
    )

    for document in sorted(
        documents,
        key=lambda item: item[
            "source"
        ],
    ):

        hasher.update(
            document[
                "source"
            ].encode(
                "utf-8"
            )
        )

        hasher.update(
            document[
                "content"
            ].encode(
                "utf-8"
            )
        )

    return hasher.hexdigest()


# =========================================================
# IN-MEMORY CACHE
# =========================================================

_embedding_cache = None
_embedding_fingerprint = None


def clear_embedding_cache():
    """
    Clear the in-memory knowledge index.
    """

    global _embedding_cache
    global _embedding_fingerprint

    _embedding_cache = None
    _embedding_fingerprint = None


# =========================================================
# KNOWLEDGE CHANGE DETECTION
# =========================================================

def knowledge_has_changed() -> bool:

    current_fingerprint = (
        calculate_knowledge_fingerprint()
    )

    if (
        _embedding_fingerprint
        is None
    ):

        return True

    return (
        current_fingerprint
        != _embedding_fingerprint
    )


# =========================================================
# BUILD EMBEDDING CACHE
# =========================================================

def build_embedding_cache() -> list[dict]:

    global _embedding_cache
    global _embedding_fingerprint

    chunks = (
        create_knowledge_chunks()
    )

    cached_chunks = []

    for chunk in chunks:

        embedding = embed_text(
            chunk[
                "content"
            ]
        )

        cached_chunks.append(
            {
                **chunk,
                "embedding": (
                    embedding
                ),
            }
        )

    _embedding_cache = (
        cached_chunks
    )

    _embedding_fingerprint = (
        calculate_knowledge_fingerprint()
    )

    logger.info(
        "Embedding cache built chunks=%s",
        len(cached_chunks),
    )

    return _embedding_cache


# =========================================================
# SAVE EMBEDDING INDEX
# =========================================================

def save_embedding_index() -> None:

    global _embedding_cache
    global _embedding_fingerprint

    if _embedding_cache is None:

        build_embedding_cache()

    INDEX_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    payload = {
        "fingerprint": (
            _embedding_fingerprint
        ),
        "chunks": (
            _embedding_cache
        ),
    }

    with INDEX_PATH.open(
        "wb"
    ) as file:

        pickle.dump(
            payload,
            file,
        )

    logger.info(
        "Embedding index saved: %s",
        INDEX_PATH,
    )


# =========================================================
# LOAD EMBEDDING INDEX
# =========================================================

def load_embedding_index() -> bool:

    global _embedding_cache
    global _embedding_fingerprint

    if not INDEX_PATH.exists():

        return False

    try:

        with INDEX_PATH.open(
            "rb"
        ) as file:

            payload = pickle.load(
                file
            )

        fingerprint = payload.get(
            "fingerprint"
        )

        chunks = payload.get(
            "chunks"
        )

        if fingerprint is None:

            return False

        if chunks is None:

            return False

        if not isinstance(
            chunks,
            list,
        ):

            return False

        _embedding_cache = (
            chunks
        )

        _embedding_fingerprint = (
            fingerprint
        )

        return True

    except Exception as exc:

        logger.warning(
            "Unable to load embedding index: %s",
            exc,
        )

        clear_embedding_cache()

        return False


# =========================================================
# BUILD OR LOAD INDEX
# =========================================================

def build_or_load_index() -> list[dict]:

    global _embedding_cache
    global _embedding_fingerprint

    current_fingerprint = (
        calculate_knowledge_fingerprint()
    )

    # -----------------------------------------------------
    # 1. Reuse valid in-memory cache
    # -----------------------------------------------------

    if (
        _embedding_cache
        is not None
        and
        _embedding_fingerprint
        == current_fingerprint
    ):

        return _embedding_cache

    # -----------------------------------------------------
    # 2. Try persistent disk index
    # -----------------------------------------------------

    loaded = (
        load_embedding_index()
    )

    if (
        loaded
        and
        _embedding_fingerprint
        == current_fingerprint
    ):

        return _embedding_cache

    # -----------------------------------------------------
    # 3. Knowledge changed or index missing/stale
    # -----------------------------------------------------

    logger.info(
        "Knowledge changed or stale index detected. "
        "Rebuilding index."
    )

    build_embedding_cache()

    # -----------------------------------------------------
    # 4. Save refreshed index
    # -----------------------------------------------------

    save_embedding_index()

    return _embedding_cache


# =========================================================
# CACHE ACCESS
# =========================================================

def get_embedding_cache() -> list[dict]:

    return build_or_load_index()


# =========================================================
# SEMANTIC RETRIEVAL
# =========================================================

def retrieve_knowledge(
    query: str,
    max_results: int = 3,
    minimum_score: float = 0.25,
) -> list[dict]:

    if max_results <= 0:

        return []

    cleaned_query = (
        query.strip()
    )

    if not cleaned_query:

        return []

    cached_chunks = (
        build_or_load_index()
    )

    if not cached_chunks:

        return []

    query_embedding = embed_text(
        cleaned_query
    )

    scored_chunks = []

    for chunk in cached_chunks:

        score = cosine_similarity(
            query_embedding,
            chunk[
                "embedding"
            ],
        )

        if score >= minimum_score:

            scored_chunks.append(
                (
                    score,
                    chunk,
                )
            )

    scored_chunks.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    results = [
        {
            "source": (
                chunk[
                    "source"
                ]
            ),
            "chunk_id": (
                chunk[
                    "chunk_id"
                ]
            ),
            "chunk_index": (
                chunk[
                    "chunk_index"
                ]
            ),
            "content": (
                chunk[
                    "content"
                ]
            ),
            "score": round(
                score,
                4,
            ),
        }
        for (
            score,
            chunk,
        )
        in scored_chunks[
            :max_results
        ]
    ]

    logger.info(
        "Knowledge retrieval query completed results=%s sources=%s",
        len(results),
        list(
            dict.fromkeys(
                result[
                    "source"
                ]
                for result
                in results
            )
        ),
    )

    return results